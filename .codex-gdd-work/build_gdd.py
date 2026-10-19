from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\alons\FLOWSTATE")
SOURCE = Path(r"C:\Users\alons\Downloads\FlowState_GDD.docx")
OUTPUT = ROOT / "Deliverables" / "FlowState_GDD_Actualizado_2026-10-06.docx"

NAVY = "111827"
INK = "101725"
CHARCOAL = "1F2937"
MID = "667085"
PALE = "F2F5F9"
WHITE = "FFFFFF"
BORDER = "303744"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_fixed(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_run(run, font="Arial", size=11, bold=False, italic=False, color=INK):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def body(doc, text="", *, bold_lead=None, keep=False, after=8):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run(r2)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def small_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run(r, font="Arial", size=9, italic=True, color=MID)
    return p


def h1(doc, number, title):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    r1 = p.add_run(f"{number}. ")
    set_run(r1, font="Arial", size=14, bold=True, color=NAVY)
    r2 = p.add_run(title)
    set_run(r2, font="Arial Black", size=14, bold=True, color=NAVY)
    return p


def h2(doc, title, page_break=False):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(title)
    set_run(r, font="Courier New", size=13, bold=True, color=CHARCOAL)
    return p


def h3(doc, title):
    p = doc.add_paragraph(style="Heading 3")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    r = p.add_run(title)
    set_run(r, font="Courier New", size=11, bold=True, color=CHARCOAL)
    return p


def table(doc, headers, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    set_table_fixed(t)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for i, width in enumerate(widths):
            for cell in t.columns[i].cells:
                cell.width = Inches(width)
    header = t.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, NAVY)
        set_cell_margins(cell)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(label))
        set_run(r, size=font_size, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = t.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, PALE if row_index % 2 else WHITE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run(r, size=font_size, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return t


def add_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def figure(doc, image_path, caption, alt, *, max_width=6.65, max_height=7.4):
    image_path = Path(image_path)
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    from PIL import Image

    with Image.open(image_path) as im:
        px_w, px_h = im.size
    aspect = px_w / px_h
    width = max_width
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    add_alt_text(shape, caption, alt)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    c.paragraph_format.keep_together = True
    r = c.add_run(caption)
    set_run(r, font="Arial", size=9, italic=True, color=MID)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break()
    p.runs[0]._element.get_or_add_br().set(qn("w:type"), "page") if False else None


def page_break(doc):
    doc.add_page_break()


def configure_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SOURCE))
    configure_update_fields(doc)

    # New work is appended in the existing section so the retained border and footer continue.
    page_break(doc)
    h1(doc, 13, "Estado actual del prototipo")
    body(
        doc,
        "Esta sección registra el estado verificable del proyecto Unity al 6 de octubre de 2026. "
        "Complementa —sin reemplazar— las definiciones de diseño de las secciones anteriores y separa lo implementado, "
        "lo validado como prototipo y lo que todavía requiere producción o pruebas.",
    )
    small_note(
        doc,
        "Criterio editorial: el código, las escenas, los assets y los informes de prueba del repositorio son la fuente de verdad. "
        "Cuando una función existe pero no está habilitada por defecto o no tiene cobertura suficiente, se declara expresamente.",
    )

    h2(doc, "13.1 Lectura del estado")
    table(
        doc,
        ["Etiqueta", "Interpretación en este documento"],
        [
            ["Implementado", "Existe en escenas, prefabs, assets o código integrado del proyecto."],
            ["Prototipo validado", "Fue ejecutado y comprobado en el alcance indicado; todavía puede requerir pulido de producción."],
            ["Infraestructura", "La base técnica existe, pero el contenido final, la activación o la integración completa siguen pendientes."],
            ["Pendiente / no validado", "No hay evidencia suficiente para considerarlo cerrado, estable o listo para entrega."],
        ],
        widths=[1.55, 5.35],
    )

    h2(doc, "13.2 Resumen ejecutivo de implementación")
    table(
        doc,
        ["Área", "Estado", "Resultado verificable"],
        [
            ["Flujo base", "Implementado", "MainMenu_FlowState abre el juego y JUGAR carga Gameplay."],
            ["Graffiti", "Implementado", "Pintura libre y decal, color HSV, cinco boquillas, recursos de aerosol y HUD."],
            ["Movimiento", "Implementado", "Tercera persona, sprint, salto, cámara, vault, trepar, wall jump y air roll."],
            ["Menú principal", "Prototipo validado", "Menú diegético con galería, configuración, salida y transición de pintura."],
            ["Introducción", "Prototipo validado", "Secuencia de seis ventanas verticales; guardada desactivada por defecto."],
            ["FSSRS", "Implementado", "Render estilizado URP con contornos, tramas, posterización, paletas y estados emocionales."],
            ["Cinemáticas", "Infraestructura", "Timeline, pintado modular, props animados y overlays; falta cierre editorial y sonoro."],
            ["Audio", "Prototipo", "Dispositivo musical diegético operativo; audio de spray silenciado por defecto."],
            ["QA de build", "Pendiente", "No hay build Windows empaquetada, benchmark de frame time ni prueba prolongada con control físico."],
        ],
        widths=[1.28, 1.48, 4.12],
        font_size=8.5,
    )

    page_break(doc)
    h2(doc, "13.3 Flujo jugable y arquitectura de escenas")
    body(
        doc,
        "El ejecutable está organizado alrededor de dos escenas incluidas en Build Settings. La entrada es el menú principal y la transición "
        "de JUGAR conduce a la escena de gameplay. Los sistemas persistentes —como reputación y preferencias— se apoyan en PlayerPrefs; "
        "los sistemas de interacción y presentación viven en componentes especializados dentro de cada escena.",
    )
    table(
        doc,
        ["Orden", "Escena", "Función actual"],
        [
            ["0", "Assets/02_Escenas/MainMenu_FlowState.unity", "Entrada, navegación, galería, configuración, introducción opcional y transición."],
            ["1", "Assets/02_Escenas/Gameplay.unity", "Movimiento, cámara, graffiti, recursos, HUD, reputación, música y pruebas de traversal."],
        ],
        widths=[0.7, 3.45, 2.75],
        font_size=8.5,
    )
    h3(doc, "Flujo mínimo verificable")
    table(
        doc,
        ["Paso", "Acción", "Respuesta esperada"],
        [
            ["1", "Iniciar el proyecto", "Aparece MainMenu_FlowState como primera escena de build."],
            ["2", "Elegir JUGAR", "Se reproduce la transición violeta orgánica durante aproximadamente 1,05 s."],
            ["3", "Entrar a Gameplay", "Se habilitan locomoción, cámara y sistemas jugables integrados."],
            ["4", "Activar graffiti con Q", "La cámara pasa a primera persona y abre el flujo de pintura."],
            ["5", "Salir con Esc", "El modo graffiti se cierra y vuelve el control normal, según el contexto activo."],
        ],
        widths=[0.65, 2.45, 3.80],
        font_size=8.5,
    )

    h2(doc, "13.4 Controles actuales")
    table(
        doc,
        ["Contexto", "Entrada", "Acción"],
        [
            ["Gameplay", "WASD / stick", "Desplazamiento del personaje."],
            ["Gameplay", "Mouse / stick derecho", "Control de cámara."],
            ["Gameplay", "Shift", "Sprint."],
            ["Gameplay", "Espacio", "Salto y acciones contextuales de traversal."],
            ["Gameplay", "Q", "Entrar en modo graffiti."],
            ["Graffiti", "Mouse", "Apuntar y pintar sobre superficies habilitadas."],
            ["Graffiti", "Tab", "Recorrer las cinco boquillas."],
            ["Graffiti", "R", "Agitar la lata y recuperar la mezcla."],
            ["Graffiti", "F", "Cambiar la lata; el suministro es infinito en el prototipo actual."],
            ["Graffiti", "Esc", "Salir del modo graffiti."],
            ["Combate", "Botón derecho", "Bloqueo de objetivo si hay un candidato válido."],
            ["Música", "M / P", "Alternar dispositivo musical / reproducir."],
            ["Música", "Flechas", "Ajustar BPM."],
            ["Música", "Teclas numéricas", "Seleccionar pistas disponibles."],
            ["Menú", "Flechas, A/D, d-pad o stick", "Cambiar opción."],
            ["Menú", "Enter, botón sur o clic", "Confirmar."],
            ["Menú", "Esc o botón este", "Volver."],
            ["Menú", "Rueda del mouse", "Navegar entre opciones."],
        ],
        widths=[1.25, 2.15, 3.50],
        font_size=8.2,
    )
    small_note(doc, "El proyecto tiene instalado Input System 1.13.1 y acepta ambos manejadores; el gameplay inspeccionado utiliza principalmente UnityEngine.Input heredado.")

    page_break(doc)
    h2(doc, "13.5 Sistema de graffiti")
    body(
        doc,
        "El graffiti ya funciona como una herramienta jugable integrada: permite pintura libre y colocación de decals, selección cromática HSV, "
        "cambio de boquilla, consumo de recursos y visualización mediante HUD. La superficie de pintura mantiene un límite de 12.000 stamps "
        "por canvas para acotar memoria y costo de dibujo.",
    )
    table(
        doc,
        ["Boquilla", "Perfil interno", "Lectura visual y uso"],
        [
            ["Pequeña", "Needle", "Trazo fino y preciso; apropiado para detalle y contorno."],
            ["Difuminar", "Soft", "Aplicación suave, difusa y de transición."],
            ["Grande", "FatCap", "Cobertura amplia y rápida."],
            ["Trazo", "Chisel", "Marca direccional, plana y caligráfica."],
            ["Mediana", "Splatter", "Dispersión irregular y gestual."],
        ],
        widths=[1.35, 1.45, 4.10],
        font_size=8.8,
    )
    figure(
        doc,
        ROOT / "Assets/Screenshots/GraffitiMenu_Refined_Close.png",
        "Figura 13.1 — Neo Spray Lab: selector diegético de boquillas dentro del espacio jugable.",
        "Captura del selector Neo Spray Lab con cinco boquillas montadas en un soporte tridimensional frente al muro de graffiti.",
    )

    page_break(doc)
    h2(doc, "13.6 Recursos del aerosol y HUD")
    body(
        doc,
        "La lata simula dos recursos separados, ambos expresados entre 0 y 100. La pintura representa contenido disponible; la mezcla representa "
        "la presión y homogeneidad. Cuando la mezcla baja, el trazo pierde opacidad, radio y densidad y puede presentar interrupciones. "
        "Cada boquilla aplica multiplicadores propios de consumo.",
    )
    table(
        doc,
        ["Parámetro", "Valor actual", "Comportamiento"],
        [
            ["Consumo de pintura", "4 por segundo", "Disminuye mientras se pinta."],
            ["Consumo de mezcla al pintar", "7,5 por segundo", "Acelera la inestabilidad del spray."],
            ["Pérdida de mezcla en reposo", "0,12 por segundo", "Evita que la mezcla permanezca estable indefinidamente."],
            ["Pintura crítica", "20", "Activa lectura de advertencia en el sistema y el HUD."],
            ["Mezcla inestable", "35", "Introduce degradación perceptible del trazo."],
            ["Agitar con R", "1,25 s", "Restaura la mezcla a 100."],
            ["Cambiar lata con F", "0,82 s", "Repone la lata; suministro infinito habilitado en el prototipo."],
        ],
        widths=[2.20, 1.45, 3.25],
        font_size=8.5,
    )
    figure(
        doc,
        ROOT / "Assets/Screenshots/spray_hud_full_final.png",
        "Figura 13.2 — HUD de aerosol con barras de pintura y mezcla durante gameplay.",
        "Captura de juego en tercera persona con el HUD de aerosol mostrando sus indicadores de pintura y mezcla.",
    )

    page_break(doc)
    h2(doc, "13.7 Movimiento, cámara, parkour y combate")
    body(
        doc,
        "El controlador de personaje reúne locomoción en tercera persona, sprint, salto con coyote time y buffer, colisión de cámara y ajuste de FOV. "
        "El modo graffiti conserva desplazamiento reducido y cambia a cámara en primera persona. El traversal incorpora vault bajo y medio, trepar, "
        "top-out, escalada continua, wall jump y air roll.",
    )
    table(
        doc,
        ["Subsistema", "Estado", "Notas de integración"],
        [
            ["Locomoción base", "Implementado", "Caminar, sprint, salto, gravedad y respuesta de cámara."],
            ["Asistencia de salto", "Implementado", "Coyote time y buffer para reducir fallos de intención."],
            ["Vault", "Implementado", "Variantes baja y media según obstáculo."],
            ["Climb / top-out", "Implementado", "Ascenso continuo y recuperación sobre borde."],
            ["Wall jump / air roll", "Implementado", "Movilidad aérea y recuperación."],
            ["Cámara", "Implementado", "Colisión, zoom/FOV y transición de modos."],
            ["Lock-on", "Implementado", "Adquisición hasta 18 m y retención hasta 24 m; se bloquea en graffiti y música."],
            ["Combate completo", "Infraestructura", "Existe selección/resaltado de objetivos; no se declara un loop de combate final."],
            ["Cobertura de regresión", "Pendiente", "No hay una suite dedicada que cubra todo el parkour de extremo a extremo."],
        ],
        widths=[1.65, 1.35, 3.90],
        font_size=8.4,
    )
    figure(
        doc,
        ROOT / "Assets/Screenshots/third_person_camera_final.png",
        "Figura 13.3 — Cámara en tercera persona integrada en la escena de gameplay.",
        "Captura del personaje en tercera persona dentro del entorno de prueba urbano.",
    )

    page_break(doc)
    h2(doc, "13.8 FSSRS y dirección visual implementada")
    body(
        doc,
        "Flow State Stylized Rendering System (FSSRS) es una Renderer Feature personalizada para URP acompañada por un volumen de control. "
        "Combina contornos, posterización, motas de tinta, halftone, hatching, influencia de paleta, elevación de papel, saturación y acentos. "
        "Las paletas emocionales verificadas son Neutral, Duda, Rabia, Flujo creativo y Claridad.",
    )
    table(
        doc,
        ["Capa visual", "Función expresiva"],
        [
            ["Contorno y ecos", "Separa siluetas y añade bordes de lenguaje cómic."],
            ["Posterización", "Reduce gradientes y refuerza masas gráficas."],
            ["Halftone y hatching", "Introduce trama impresa, volumen y tensión."],
            ["Motas de tinta", "Rompe la limpieza digital y suma materialidad."],
            ["Paleta emocional", "Desplaza color y contraste según el estado narrativo."],
            ["Paper lift", "Evoca capas de papel y collage urbano."],
        ],
        widths=[2.15, 4.75],
        font_size=8.8,
    )
    figure(
        doc,
        ROOT / "Assets/Screenshots/fssrs-v3-state-flow-max-rainbow-a.png",
        "Figura 13.4 — FSSRS en estado de Flujo creativo con paleta de máxima intensidad.",
        "Retrato estilizado del personaje con contornos, trama, hatching y una paleta multicolor de flujo creativo.",
    )

    page_break(doc)
    h2(doc, "13.9 Menú principal vivo")
    body(
        doc,
        "El menú principal es una escena diegética, no una pantalla plana. Presenta JUGAR, GALERÍA, CONFIG y SALIR; admite teclado, mouse y gamepad, "
        "y conserva una opción de movimiento reducido. Galería y configuración se resuelven con paneles world-space de papel y tinta. "
        "La entrada a gameplay usa una transición de pintura violeta suave y orgánica.",
    )
    table(
        doc,
        ["Opción", "Estado", "Resultado actual"],
        [
            ["JUGAR", "Implementado", "Carga Gameplay después de la transición."],
            ["GALERÍA", "Implementado", "Panel diegético para contenido visual."],
            ["CONFIG", "Implementado", "Panel world-space con preferencias y movimiento reducido."],
            ["SALIR", "Implementado", "Cierre de la aplicación en build."],
        ],
        widths=[1.25, 1.45, 4.20],
        font_size=8.8,
    )
    figure(
        doc,
        ROOT / "Docs/AI/MainMenuPreview.png",
        "Figura 13.5 — Menú principal de Flow State con composición urbana, stencil y lata interactiva.",
        "Captura del menú principal de Flow State con muro ilustrado, logotipo, opciones y una lata de aerosol en primer plano.",
    )

    page_break(doc)
    h2(doc, "13.10 Introducción y presentación narrativa")
    body(
        doc,
        "La infraestructura de introducción está integrada al menú. El prototipo revisado dura aproximadamente 5,10 s y utiliza seis ventanas "
        "verticales diseñadas como dioramas sobre fondo negro. Emplea una cámara de contenido, una Render Texture de contenido y dos máscaras. "
        "La respuesta reactiva de cámara fue corregida para respetar movimiento reducido con una escala de 0,6.",
    )
    table(
        doc,
        ["Elemento", "Estado actual"],
        [
            ["Composición de seis ventanas", "Prototipo completo y revisado."],
            ["IntroMode", "Guardado como Disabled en la escena principal."],
            ["Movimiento reducido", "Integrado y corregido en la respuesta de cámara."],
            ["Audio final", "Pendiente."],
            ["Handoff automático definitivo", "Pendiente de decisión e integración final."],
        ],
        widths=[2.65, 4.25],
        font_size=8.8,
    )
    figure(
        doc,
        ROOT / "Captures/Phase3ABReview/IntroPanels_02_SixComplete.png",
        "Figura 13.6 — Prototipo de introducción con las seis ventanas narrativas completas.",
        "Seis paneles verticales iluminados sobre fondo negro que componen la introducción prototipo de Flow State.",
    )

    page_break(doc)
    h2(doc, "13.11 Cinemáticas y material de presentación")
    body(
        doc,
        "El proyecto contiene una base de producción cinematográfica con Timeline, un sistema de pintado modular, animación de lata y personaje, "
        "y overlays de paleta, emoción y logotipo. El asset FlowState_Trailer_25s_98s.playable organiza una pieza temporal de presentación; "
        "su existencia se considera infraestructura y prototipo, no máster final exportado.",
    )
    figure(
        doc,
        ROOT / "Captures/FlowState_v6_FSSRS_contact.png",
        "Figura 13.7 — Hoja de contacto de la presentación cinematográfica con variaciones FSSRS.",
        "Hoja panorámica de contacto con múltiples encuadres del personaje, el aerosol, la ciudad y tratamientos gráficos de Flow State.",
        max_width=6.65,
        max_height=6.0,
    )
    table(
        doc,
        ["Pieza técnica", "Ubicación / función"],
        [
            ["FlowStateCinematicShowcase.cs", "Orquesta el showcase y sus estados visuales."],
            ["FlowStateCinematicPainter.cs", "Controla revelados y trazos de presentación."],
            ["FlowState_Trailer_25s_98s.playable", "Timeline de la pieza prototipo."],
            ["Overlays", "Conectan logotipo, emoción y paleta con el montaje."],
        ],
        widths=[2.70, 4.20],
        font_size=8.5,
    )

    page_break(doc)
    h2(doc, "13.12 Sistemas auxiliares integrados")
    table(
        doc,
        ["Sistema", "Estado", "Descripción actual"],
        [
            ["Reputación", "Implementado", "Recognition, CrewRespect y CityImpact; total ponderado 35/30/35 y persistencia en PlayerPrefs."],
            ["Música diegética", "Prototipo", "Dispositivo con reproducción, selección de pistas y BPM; valor inicial observado: 92."],
            ["Teletransporte de alcantarillas", "Implementado", "Red con radio de 3,5 y fade 0,3 / negro 0,12 / fade 0,4 s."],
            ["Pausa", "Implementado", "Control de pausa integrado en el proyecto."],
            ["Transición de cámara", "Implementado", "Switcher y transiciones para modos y dispositivos."],
            ["Targeting", "Implementado", "Selección y resaltado de objetivo compatibles con el lock-on."],
            ["Audio de spray", "Infraestructura", "Fallback procedural presente; salida silenciada por defecto."],
        ],
        widths=[1.65, 1.30, 3.95],
        font_size=8.4,
    )
    figure(
        doc,
        ROOT / "Assets/Screenshots/music_device_prototype_final.png",
        "Figura 13.8 — Prototipo del dispositivo musical diegético dentro de gameplay.",
        "Captura de un dispositivo musical interactivo en el entorno de prueba de Flow State.",
    )

    page_break(doc)
    h2(doc, "13.13 Base técnica y fuentes de verdad")
    table(
        doc,
        ["Capa", "Versión / ubicación", "Responsabilidad"],
        [
            ["Motor", "Unity 6000.0.41f1", "Runtime y editor del proyecto."],
            ["Render", "URP 17.0.4", "Pipeline base de iluminación y postproceso."],
            ["Input", "Input System 1.13.1 + legado", "Compatibilidad de dispositivos; lógica principal inspeccionada usa UnityEngine.Input."],
            ["Movimiento", "FPSController.cs", "Locomoción, traversal, cámara y modo graffiti."],
            ["Graffiti", "GraffitiPainter.cs", "Pintura, color, boquillas, decals e integración."],
            ["Superficie", "GraffitiSurfaceCanvas.cs", "Acumulación y límite de stamps."],
            ["Recursos", "SprayResourceSystem.cs", "Pintura, mezcla, agitado, cambio de lata y penalizaciones."],
            ["HUD", "SprayResourceHUD.cs", "Lectura visual de recursos y estados críticos."],
            ["Combate", "CombatLockOnController.cs", "Lock-on, distancias y bloqueos contextuales."],
            ["Audio", "DiegeticMusicPlayer.cs", "Dispositivo musical, BPM y pistas."],
            ["Reputación", "ReputationSystem.cs", "Ejes, total ponderado y persistencia."],
            ["Teleport", "SewerTeleportNetwork.cs", "Viaje y fundidos entre puntos."],
            ["Cinemáticas", "Assets/10_Cinematics", "Runtime, Timeline, overlays y material de captura."],
            ["Documentación", "Docs/AI", "Decisiones, fases, resultados de pruebas y estado operativo."],
        ],
        widths=[1.30, 2.55, 3.05],
        font_size=7.9,
    )
    small_note(doc, "Las rutas y nombres anteriores sirven como índice operativo para futuras actualizaciones. NotebookLM debe pedir evidencia del repositorio antes de convertir una intención de diseño en estado implementado.")

    page_break(doc)
    h2(doc, "13.14 Validación registrada")
    table(
        doc,
        ["Fecha / alcance", "Resultado", "Lectura correcta"],
        [
            ["21-09-2026 — menú e intro", "19/19 Play Mode y 7/7 Edit Mode", "Sin warnings del menú o shader en el alcance informado; intro todavía desactivada por defecto."],
            ["23-09-2026 — Neo Spray Lab", "7 Edit Mode; 19 Play Mode combinados", "Una falla dependiente del orden en migración pasó al ejecutarse aislada; no ocultar esta salvedad."],
            ["Revisión actual del repositorio", "Escenas, scripts, assets y capturas presentes", "Confirma integración documental, no sustituye una certificación de build final."],
        ],
        widths=[2.05, 1.65, 3.20],
        font_size=8.3,
    )
    h3(doc, "Pendientes de validación")
    table(
        doc,
        ["Pendiente", "Criterio de cierre propuesto"],
        [
            ["Build Windows empaquetada", "Generar build reproducible, ejecutar smoke test y registrar versión."],
            ["Rendimiento", "Medir frame time y memoria en hardware objetivo con FSSRS y graffiti activos."],
            ["Control físico", "Realizar una sesión prolongada con gamepad y revisar navegación, cámara y remapeo."],
            ["Parkour", "Crear pruebas dedicadas para vault, climb, top-out, wall jump y air roll."],
            ["Introducción", "Definir activación, audio final y handoff automático al menú o gameplay."],
            ["Audio de aerosol", "Aprobar diseño sonoro y retirar el mute por defecto cuando corresponda."],
            ["Persistencia", "Validar migración y compatibilidad de PlayerPrefs entre versiones."],
        ],
        widths=[2.15, 4.75],
        font_size=8.5,
    )

    h2(doc, "13.15 Convención de actualización para NotebookLM")
    body(
        doc,
        "Para mantener este GDD vivo sin degradar su confiabilidad, cada incorporación debe registrar una fecha, una fuente verificable y una etiqueta de estado. "
        "Las propuestas se conservan como intención hasta que exista evidencia en escenas, código, assets o pruebas. Las cifras solo se actualizan desde valores serializados, "
        "constantes del código o resultados medidos.",
    )
    table(
        doc,
        ["Regla", "Aplicación"],
        [
            ["No sobrescribir historia", "Agregar una nota fechada o una nueva versión; conservar decisiones anteriores relevantes."],
            ["Separar intención y realidad", "Usar las etiquetas Implementado, Prototipo validado, Infraestructura o Pendiente."],
            ["Citar evidencia interna", "Anotar escena, script, asset, informe o captura que respalda el cambio."],
            ["No inventar métricas", "Si no existe medición, declarar pendiente de benchmark o QA."],
            ["Actualizar controles", "Registrar contexto, entrada, acción y posibles conflictos."],
            ["Actualizar imágenes", "Usar capturas del proyecto, pie de figura y texto alternativo; retirar solo si quedan obsoletas."],
            ["Conservar diseño", "Mantener A4, borde gráfico, tipografías, jerarquía, tablas y pie de página de este documento."],
        ],
        widths=[2.10, 4.80],
        font_size=8.5,
    )
    small_note(doc, "Corte documental: 6 de octubre de 2026. Esta sección debe revisarse después de cada hito jugable, integración audiovisual o ciclo formal de QA.")

    doc.core_properties.title = "Flow State — Game Design Document"
    doc.core_properties.subject = "GDD maestro actualizado con el estado verificable del prototipo"
    doc.core_properties.keywords = "Flow State, GDD, graffiti, parkour, Unity, FSSRS"
    doc.core_properties.comments = "Secciones 1–12 preservadas; sección 13 agregada desde evidencia del proyecto."
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
