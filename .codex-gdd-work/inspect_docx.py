from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def fmt_length(value):
    return None if value is None else {"pt": value.pt, "in": value.inches}


def style_payload(style):
    font = style.font
    pf = style.paragraph_format
    return {
        "name": style.name,
        "type": str(style.type),
        "base_style": style.base_style.name if style.base_style else None,
        "font_name": font.name,
        "font_size": fmt_length(font.size),
        "bold": font.bold,
        "italic": font.italic,
        "color": str(font.color.rgb) if font.color and font.color.rgb else None,
        "alignment": str(pf.alignment),
        "space_before": fmt_length(pf.space_before),
        "space_after": fmt_length(pf.space_after),
        "line_spacing": str(pf.line_spacing),
        "keep_with_next": pf.keep_with_next,
        "keep_together": pf.keep_together,
        "page_break_before": pf.page_break_before,
        "left_indent": fmt_length(pf.left_indent),
        "first_line_indent": fmt_length(pf.first_line_indent),
    }


path = Path(sys.argv[1])
doc = Document(path)

payload = {
    "path": str(path),
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shape_count": len(doc.inline_shapes),
    "sections": [],
    "styles": {},
    "paragraphs": [],
    "tables": [],
    "package_parts": [],
}

for section in doc.sections:
    payload["sections"].append({
        "page_width": fmt_length(section.page_width),
        "page_height": fmt_length(section.page_height),
        "left_margin": fmt_length(section.left_margin),
        "right_margin": fmt_length(section.right_margin),
        "top_margin": fmt_length(section.top_margin),
        "bottom_margin": fmt_length(section.bottom_margin),
        "header_distance": fmt_length(section.header_distance),
        "footer_distance": fmt_length(section.footer_distance),
        "different_first_page": section.different_first_page_header_footer,
        "start_type": str(section.start_type),
    })

for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption", "Table Grid", "TOC Heading", "toc 1", "toc 2", "toc 3"]:
    try:
        payload["styles"][name] = style_payload(doc.styles[name])
    except KeyError:
        pass

for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text or paragraph.style.name.startswith("Heading"):
        p = {
            "index": index,
            "style": paragraph.style.name,
            "text": text,
            "alignment": str(paragraph.alignment),
            "runs": [],
        }
        for run in paragraph.runs[:8]:
            rpr = run._element.rPr
            fonts = None
            if rpr is not None and rpr.rFonts is not None:
                fonts = {
                    "ascii": rpr.rFonts.get(qn("w:ascii")),
                    "hAnsi": rpr.rFonts.get(qn("w:hAnsi")),
                }
            p["runs"].append({
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "font": run.font.name,
                "size": fmt_length(run.font.size),
                "color": str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None,
                "fonts": fonts,
            })
        payload["paragraphs"].append(p)

for ti, table in enumerate(doc.tables):
    grid = table._tbl.tblGrid
    payload["tables"].append({
        "index": ti,
        "style": table.style.name if table.style else None,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "grid_cols": [col.get(qn("w:w")) for col in grid.gridCol_lst] if grid is not None else [],
        "first_row": [cell.text for cell in table.rows[0].cells] if table.rows else [],
    })

with zipfile.ZipFile(path) as archive:
    for info in archive.infolist():
        data = archive.read(info.filename)
        payload["package_parts"].append({
            "path": info.filename,
            "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
        })

print(json.dumps(payload, ensure_ascii=False, indent=2))
