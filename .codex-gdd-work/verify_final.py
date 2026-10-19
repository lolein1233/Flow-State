from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SOURCE = Path(r"C:\Users\alons\Downloads\FlowState_GDD.docx")
FINAL = Path(r"C:\Users\alons\FLOWSTATE\Deliverables\FlowState_GDD_Actualizado_2026-10-06.docx")
PDF = Path(r"C:\Users\alons\FLOWSTATE\.codex-gdd-work\final-render\FlowState_GDD_Actualizado_2026-10-06.pdf")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def table_text(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest().upper()
            for name in names
        }


source_doc = Document(SOURCE)
final_doc = Document(FINAL)

source_paragraphs = [paragraph.text for paragraph in source_doc.paragraphs]
final_original_paragraphs = [paragraph.text for paragraph in final_doc.paragraphs[: len(source_paragraphs)]]

source_tables = [table_text(table) for table in source_doc.tables]
final_original_tables = [table_text(table) for table in final_doc.tables[: len(source_tables)]]

source_media = media_hashes(SOURCE)
final_media = media_hashes(FINAL)
source_media_digests = sorted(source_media.values())
final_media_digests = sorted(final_media.values())
missing_source_media = sorted(set(source_media_digests) - set(final_media_digests))

with zipfile.ZipFile(FINAL) as archive:
    document_xml = archive.read("word/document.xml")
    drawing_count = document_xml.count(b"<w:drawing") + document_xml.count(b"<w:pict")

report = {
    "source_sha256": sha256(SOURCE),
    "source_paragraph_count": len(source_paragraphs),
    "final_paragraph_count": len(final_doc.paragraphs),
    "original_paragraphs_match": source_paragraphs == final_original_paragraphs,
    "source_table_count": len(source_tables),
    "final_table_count": len(final_doc.tables),
    "original_tables_match": source_tables == final_original_tables,
    "source_media_count": len(source_media),
    "final_media_count": len(final_media),
    "original_media_preserved_by_content": not missing_source_media,
    "missing_source_media_hashes": missing_source_media,
    "added_media_count": len(final_media) - len(source_media),
    "final_drawing_count": drawing_count,
    "pdf_page_count": len(PdfReader(str(PDF)).pages),
    "required_section_present": any(
        paragraph.text.startswith("13.") and "PROTOTIPO" in paragraph.text.upper()
        for paragraph in final_doc.paragraphs
    ),
    "new_section_headings": [
        paragraph.text
        for paragraph in final_doc.paragraphs[len(source_paragraphs):]
        if paragraph.style.name.startswith("Heading")
    ],
    "final_exists": FINAL.exists(),
    "final_size_bytes": FINAL.stat().st_size,
}

print(json.dumps(report, ensure_ascii=False, indent=2))
