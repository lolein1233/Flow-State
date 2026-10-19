from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\alons\Downloads\DJD - Plantilla Memoria.docx")
OUT = Path(r"C:\Users\alons\FLOWSTATE\.codex-memoria-work\template-inspection.json")


def fmt_length(value):
    return None if value is None else {
        "emu": int(value),
        "pt": round(value.pt, 3),
        "in": round(value.inches, 4),
    }


def run_info(run):
    fonts = run._element.rPr.rFonts if run._element.rPr is not None else None
    return {
        "text": run.text,
        "font_name": run.font.name,
        "ascii_font": None if fonts is None else fonts.get(qn("w:ascii")),
        "hansi_font": None if fonts is None else fonts.get(qn("w:hAnsi")),
        "size_pt": None if run.font.size is None else run.font.size.pt,
        "bold": run.bold,
        "italic": run.italic,
        "color": None if run.font.color.rgb is None else str(run.font.color.rgb),
    }


doc = Document(SOURCE)
paragraphs = []
for index, paragraph in enumerate(doc.paragraphs):
    if not paragraph.text.strip() and not paragraph.runs:
        continue
    pf = paragraph.paragraph_format
    paragraphs.append({
        "index": index,
        "style": paragraph.style.name,
        "text": paragraph.text,
        "alignment": None if paragraph.alignment is None else int(paragraph.alignment),
        "before": fmt_length(pf.space_before),
        "after": fmt_length(pf.space_after),
        "line_spacing": None if pf.line_spacing is None else str(pf.line_spacing),
        "left_indent": fmt_length(pf.left_indent),
        "right_indent": fmt_length(pf.right_indent),
        "first_line_indent": fmt_length(pf.first_line_indent),
        "page_break_before": pf.page_break_before,
        "keep_with_next": pf.keep_with_next,
        "runs": [run_info(run) for run in paragraph.runs],
    })

styles = {}
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[name]
    pf = style.paragraph_format
    styles[name] = {
        "font_name": style.font.name,
        "size_pt": None if style.font.size is None else style.font.size.pt,
        "bold": style.font.bold,
        "italic": style.font.italic,
        "color": None if style.font.color.rgb is None else str(style.font.color.rgb),
        "alignment": None if pf.alignment is None else int(pf.alignment),
        "before": fmt_length(pf.space_before),
        "after": fmt_length(pf.space_after),
        "line_spacing": None if pf.line_spacing is None else str(pf.line_spacing),
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
    }

sections = []
for section in doc.sections:
    sections.append({
        "width": fmt_length(section.page_width),
        "height": fmt_length(section.page_height),
        "top": fmt_length(section.top_margin),
        "bottom": fmt_length(section.bottom_margin),
        "left": fmt_length(section.left_margin),
        "right": fmt_length(section.right_margin),
        "header_distance": fmt_length(section.header_distance),
        "footer_distance": fmt_length(section.footer_distance),
    })

parts = {}
with zipfile.ZipFile(SOURCE) as archive:
    for info in archive.infolist():
        data = archive.read(info.filename)
        parts[info.filename] = {
            "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest().upper(),
        }

payload = {
    "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest().upper(),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "section_count": len(doc.sections),
    "sections": sections,
    "styles": styles,
    "paragraphs": paragraphs,
    "parts": parts,
}
OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
