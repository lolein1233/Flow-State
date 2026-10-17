from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\alons\Downloads\DJD - Plantilla Memoria.docx")
OUTPUT = Path(r"C:\Users\alons\FLOWSTATE\Deliverables\FlowState_Memoria_Proyecto_Titulo_2026-10-06.docx")
PDF = Path(r"C:\Users\alons\FLOWSTATE\.codex-memoria-work\final-render-v2\FlowState_Memoria_Proyecto_Titulo_2026-10-06.pdf")
EXPECTED_SOURCE_SHA256 = "BE8B8A9225062BE01D79AC11CF0511C17A9B9CEE26D2DFE51BABAD547EFF5127"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        return {
            name: sha256_bytes(zf.read(name))
            for name in zf.namelist()
            if name.startswith("word/media/")
        }


def section_words(paragraphs, start_prefix: str, end_prefix: str) -> int:
    active = False
    text = []
    for p in paragraphs:
        t = p.text.strip()
        is_h1 = bool(p.style and p.style.name == "Heading 1")
        if is_h1 and t.startswith(start_prefix):
            active = True
        elif active and is_h1 and t.startswith(end_prefix):
            break
        elif active:
            text.append(t)
    return len(re.findall(r"\b[\wÁÉÍÓÚÜÑáéíóúüñ]+\b", " ".join(text)))


def main() -> None:
    assert SOURCE.exists() and OUTPUT.exists() and PDF.exists()
    assert sha256_file(SOURCE) == EXPECTED_SOURCE_SHA256

    source_doc = Document(SOURCE)
    output_doc = Document(OUTPUT)
    text = "\n".join(p.text for p in output_doc.paragraphs)
    headings = [
        p.text.strip()
        for p in output_doc.paragraphs
        if p.style and p.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]
    expected_h1 = [
        "1.- Índice",
        "2.- Agradecimientos",
        "3.- Introducción",
        "4.- Marco Investigativo–Conceptual",
        "5.- Estado del Arte",
        "6.- Bitácora de Proyecto",
        "7.- Conclusiones del Desarrollo del Proyecto",
        "8.- Bibliografía y Ludografía",
        "9.- Anexos y Estado de Cierre",
    ]
    for heading in expected_h1:
        assert heading in headings, heading

    forbidden = [
        "Nombre del proyecto",
        "Nombre Estudiantes",
        "Escribir en tercera persona",
        "En este capítulo se debe",
        "Debe contener entre",
        "Texto de ejemplo",
    ]
    for phrase in forbidden:
        assert phrase not in text, phrase
    assert "�" not in text
    assert text.count("Figura ") == 12

    source_media = media_hashes(SOURCE)
    output_media = media_hashes(OUTPUT)
    assert set(source_media.values()) <= set(output_media.values())

    with zipfile.ZipFile(OUTPUT) as zf:
        document_xml = zf.read("word/document.xml")
    assert b'TOC \\o "1-3"' in document_xml
    # The template logo may not carry alt text; every newly inserted raster figure must.
    descr_count = len(re.findall(rb'<wp:docPr[^>]+descr="[^"]+"', document_xml))
    assert descr_count >= 12, descr_count

    assert len(output_doc.sections) == len(source_doc.sections) == 1
    src_sec = source_doc.sections[0]
    out_sec = output_doc.sections[0]
    for attr in ("page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin"):
        assert getattr(src_sec, attr) == getattr(out_sec, attr), attr

    marco_words = section_words(output_doc.paragraphs, "4.- ", "5.- ")
    estado_words = section_words(output_doc.paragraphs, "5.- ", "6.- ")
    assert 1500 <= marco_words <= 2600, marco_words
    assert 1000 <= estado_words <= 1600, estado_words

    pdf = pdfium.PdfDocument(str(PDF))
    page_count = len(pdf)
    assert page_count == 32, page_count

    report = {
        "source_sha256": sha256_file(SOURCE),
        "output_sha256": sha256_file(OUTPUT),
        "output_size_bytes": OUTPUT.stat().st_size,
        "pages": page_count,
        "paragraphs": len(output_doc.paragraphs),
        "tables": len(output_doc.tables),
        "inline_shapes_including_logo": len(output_doc.inline_shapes),
        "headings": len(headings),
        "figure_alt_descriptions": descr_count,
        "marco_investigativo_words": marco_words,
        "estado_del_arte_words": estado_words,
        "source_media_preserved": True,
        "source_geometry_preserved": True,
        "required_sections_present": True,
        "template_instruction_phrases_absent": True,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
